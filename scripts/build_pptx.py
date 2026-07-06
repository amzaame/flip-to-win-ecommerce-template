#!/usr/bin/env python3
"""Génère le pitch deck Seasonly × GEO en PowerPoint modifiable (.pptx).

Charte Seasonly : écru #F7F3EC, vert sapin #1E3A2B, pêche #E8875D, sauge #9DB29A.
Usage : python3 scripts/build_pptx.py  →  pitch-deck-seasonly.pptx
"""
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN

CREAM = RGBColor(0xF7, 0xF3, 0xEC)
INK = RGBColor(0x1F, 0x2D, 0x24)
DEEP = RGBColor(0x1E, 0x3A, 0x2B)
SAGE = RGBColor(0x9D, 0xB2, 0x9A)
PEACH = RGBColor(0xE8, 0x87, 0x5D)
SAND = RGBColor(0xED, 0xE5, 0xD6)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
RED = RGBColor(0xB0, 0x43, 0x2F)

SERIF = "Georgia"
SANS = "Calibri"

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
BLANK = prs.slide_layouts[6]
W, H = prs.slide_width, prs.slide_height


def add_slide(bg=CREAM):
    s = prs.slides.add_slide(BLANK)
    s.background.fill.solid()
    s.background.fill.fore_color.rgb = bg
    return s


def box(s, x, y, w, h):
    tb = s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tb.text_frame.word_wrap = True
    return tb.text_frame


def para(tf, text, size=16, color=INK, bold=False, font=SANS, first=False,
         align=PP_ALIGN.LEFT, space_after=6, italic=False):
    p = tf.paragraphs[0] if first and not tf.paragraphs[0].runs else tf.add_paragraph()
    p.alignment = align
    p.space_after = Pt(space_after)
    r = p.add_run()
    r.text = text
    f = r.font
    f.size, f.color.rgb, f.bold, f.name, f.italic = Pt(size), color, bold, font, italic
    return p


def kicker(s, text, y=0.55):
    tb = box(s, 0.9, y, 11.5, 0.4)
    para(tb, text.upper(), size=13, color=PEACH, bold=True, first=True)


def title(s, text, y=0.95, size=36, color=DEEP, w=11.5):
    tb = box(s, 0.9, y, w, 1.1)
    para(tb, text, size=size, color=color, bold=True, font=SERIF, first=True)


def card(s, x, y, w, h, fill=WHITE):
    from pptx.enum.shapes import MSO_SHAPE
    sh = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    sh.fill.solid()
    sh.fill.fore_color.rgb = fill
    sh.line.color.rgb = RGBColor(0xE5, 0xDC, 0xCC)
    sh.line.width = Pt(0.75)
    sh.shadow.inherit = False
    sh.text_frame.word_wrap = True
    sh.text_frame.margin_left = Emu(150000)
    sh.text_frame.margin_right = Emu(150000)
    return sh


def stat_row(s, stats, y=5.6, x0=0.9, cw=2.85, gap=0.15, ch=1.35):
    for i, (big, small) in enumerate(stats):
        c = card(s, x0 + i * (cw + gap), y, cw, ch)
        tf = c.text_frame
        para(tf, big, size=26, color=DEEP, bold=True, font=SERIF, first=True, align=PP_ALIGN.CENTER, space_after=2)
        para(tf, small, size=11, color=INK, align=PP_ALIGN.CENTER)


def table(s, rows, x, y, w, col_widths=None, font_size=13, row_h=0.42):
    n_r, n_c = len(rows), len(rows[0])
    shp = s.shapes.add_table(n_r, n_c, Inches(x), Inches(y), Inches(w), Inches(row_h * n_r))
    t = shp.table
    if col_widths:
        total = sum(col_widths)
        for j, cwidth in enumerate(col_widths):
            t.columns[j].width = Emu(int(Inches(w) * cwidth / total))
    for i, row in enumerate(rows):
        for j, val in enumerate(row):
            cell = t.cell(i, j)
            cell.fill.solid()
            cell.fill.fore_color.rgb = DEEP if i == 0 else (WHITE if i % 2 else SAND)
            tfr = cell.text_frame
            tfr.word_wrap = True
            p = tfr.paragraphs[0]
            r = p.add_run()
            r.text = str(val)
            r.font.size = Pt(font_size)
            r.font.name = SANS
            r.font.bold = i == 0
            r.font.color.rgb = CREAM if i == 0 else INK
    return t


def footer(s, n):
    tb = box(s, 0.9, 7.08, 11.5, 0.35)
    para(tb, f"seasonly. × GEO — pitch deck   ·   {n}/9", size=10, color=SAGE, first=True)


# ── Slide 1 : titre ────────────────────────────────────────────────────
s = add_slide(DEEP)
tb = box(s, 0.9, 1.1, 11, 0.5)
para(tb, "PROJET SEO × GEO · JUILLET 2026", size=14, color=PEACH, bold=True, first=True)
tb = box(s, 0.9, 1.7, 11.5, 2.4)
para(tb, "Seasonly", size=60, color=CREAM, bold=True, font=SERIF, first=True, space_after=2)
para(tb, "Exister dans les réponses des IA", size=40, color=PEACH, bold=True, font=SERIF)
tb = box(s, 0.9, 4.35, 10.8, 1.3)
para(tb, "Comment une marque pionnière du facialisme, absente des SERP et des réponses de "
         "ChatGPT, Perplexity et Claude, peut devenir la source que les IA citent — "
         "en 3 actions, 18 pages et 1 robot éditorial.", size=18, color=CREAM, first=True)
tb = box(s, 0.9, 5.9, 11.5, 0.5)
para(tb, "Analyse SERP/GEO   ·   18 pages vibecodées   ·   Workflow Make   ·   KPIs", size=14, color=SAGE, first=True)

# ── Slide 2 : la marque ────────────────────────────────────────────────
s = add_slide()
kicker(s, "1 · La marque")
title(s, "Seasonly : forte en studio, invisible en ligne (hors marque)")
cards = [
    ("2018", "Fondée par Fany Péchiodat — skincare naturelle (95 %), vegan"),
    ("2", "moteurs de business : e-commerce + Skin Studios (facialisme)"),
    ("100", "studios visés en Europe d'ici 2027 (Paris, Aix, Bordeaux, Nice…)"),
    ("≈ 0 %", "de visibilité sur les requêtes génériques testées (Google + IA)"),
]
for i, (big, small) in enumerate(cards):
    c = card(s, 0.9 + i * 3.0, 2.3, 2.85, 1.8)
    tf = c.text_frame
    para(tf, big, size=32, color=DEEP, bold=True, font=SERIF, first=True, align=PP_ALIGN.CENTER, space_after=4)
    para(tf, small, size=12, color=INK, align=PP_ALIGN.CENTER)
c = card(s, 0.9, 4.6, 11.5, 1.7, fill=WHITE)
tf = c.text_frame
para(tf, "Le problème", size=15, color=PEACH, bold=True, first=True, space_after=4)
para(tf, "Les clients découvrent la catégorie (« kobido », « gua sha », « notox ») AVANT de "
         "découvrir la marque — et de plus en plus via un chatbot. Sur ces requêtes, Seasonly "
         "n'existe pas.", size=16, color=INK, italic=True)
footer(s, 2)

# ── Slide 3 : diagnostic ───────────────────────────────────────────────
s = add_slide()
kicker(s, "2 · Diagnostic GEO")
title(s, "Qui est cité à la place de Seasonly ?")
rows = [
    ["Requête testée", "Seasonly ?", "Cités à la place"],
    ["meilleur soin visage institut Paris", "ABSENTE", "Planity, Treatwell, Funbooker + instituts optimisés"],
    ["massage kobido Paris", "Partielle", "Resalib, blogs concurrents, indépendantes (130-180 €)"],
    ["meilleure marque cosmétique naturelle FR", "ABSENTE", "Blogs affiliés « Top 15 », Melvita, marques auto-publiées"],
    ["anti-âge sans injection / botox naturel", "ABSENTE", "Aroma-Zone, cliniques — alors que c'est SON territoire"],
    ["gua sha bienfaits + utilisation", "ABSENTE", "NIVEA, Mademoiselle bio… alors que Seasonly en VEND"],
]
t = table(s, rows, 0.9, 2.35, 11.5, col_widths=[3.4, 1.2, 5.6], font_size=13, row_h=0.55)
for i in (1, 3, 4, 5):
    t.cell(i, 1).text_frame.paragraphs[0].runs[0].font.color.rgb = RED
t.cell(2, 1).text_frame.paragraphs[0].runs[0].font.color.rgb = DEEP
tb = box(s, 0.9, 6.0, 11.5, 0.6)
para(tb, "1 requête sur 5 mentionne la marque — et uniquement via des articles tiers, jamais via son propre site.",
     size=16, color=INK, bold=True, first=True)
footer(s, 3)

# ── Slide 4 : patterns ─────────────────────────────────────────────────
s = add_slide()
kicker(s, "2 · Diagnostic GEO")
title(s, "Pourquoi EUX sont cités : les 7 patterns")
patterns = [
    ("📋 Listicles chiffrés", "« Top 20 », « 5 meilleures adresses » : réponse pré-mâchée"),
    ("🔢 Données précises", "Prix, notes (4,9/5), volumes (2 000 avis) = faits citables"),
    ("🧱 Structure extractible", "H2/H3, étapes, tableaux, FAQ : chunks parfaits pour le RAG"),
    ("🎓 E-E-A-T", "« avis dermatologue », auteur identifié, « on a testé »"),
    ("📅 Fraîcheur affichée", "Millésime 2026 dans le titre, date de mise à jour"),
    ("🏛️ Autorité d'agrégateur", "Planity/Treatwell : avis structurés + schema LocalBusiness"),
    ("🎯 Réponse en 1er ¶", "La définition en tête de page devient LE snippet cité"),
]
for i, (h, txt) in enumerate(patterns):
    x = 0.9 + (i % 4) * 3.0
    y = 2.3 + (i // 4) * 1.75
    c = card(s, x, y, 2.85, 1.6)
    tf = c.text_frame
    para(tf, h, size=14, color=DEEP, bold=True, font=SERIF, first=True, space_after=3)
    para(tf, txt, size=11, color=INK)
c = card(s, 9.9, 4.05, 2.5, 1.6, fill=DEEP)
tf = c.text_frame
para(tf, "Conclusion", size=14, color=PEACH, bold=True, font=SERIF, first=True, space_after=3)
para(tf, "Les IA citent le plus facile à extraire — et accessible aux crawlers.", size=11, color=CREAM)
footer(s, 4)

# ── Slide 5 : stratégie ────────────────────────────────────────────────
s = add_slide()
kicker(s, "3 · Stratégie")
title(s, "3 actions pour entrer dans les réponses")
actions = [
    ("1 · OWNED — Hub de contenu citable",
     "18 pages guides appliquant les 7 patterns : réponse directe, chiffres, FAQ, JSON-LD "
     "(FAQPage, HowTo, LocalBusiness), date et auteur. → démo Step 2"),
    ("2 · EARNED — Autorité tierce",
     "Fiches Planity/Treatwell/Google Business optimisées (viser 4,8+/5), présence dans 10 "
     "listicles « meilleur kobido/facialiste », RP presse beauté."),
    ("3 · TECH — Infrastructure GEO",
     "robots.txt ouvert aux bots IA + llms.txt + schema.org partout + flux de fraîcheur "
     "automatisé. → démo Step 3"),
]
for i, (h, txt) in enumerate(actions):
    c = card(s, 0.9, 2.25 + i * 1.15, 11.5, 1.05)
    tf = c.text_frame
    para(tf, h, size=15, color=DEEP, bold=True, font=SERIF, first=True, space_after=2)
    para(tf, txt, size=12, color=INK)
c = card(s, 0.9, 5.85, 11.5, 1.15, fill=SAND)
tf = c.text_frame
para(tf, "L'enjeu robots.txt", size=13, color=PEACH, bold=True, first=True, space_after=2)
para(tf, "Bloquer GPTBot, ClaudeBot, PerplexityBot, Google-Extended protège les médias — mais condamne "
         "une PME à l'invisibilité. Pour Seasonly : être crawlé = être cité = exister. On ouvre tout, "
         "sauf /cart, /checkout, /account.", size=12, color=INK, italic=True)
footer(s, 5)

# ── Slide 6 : les pages ────────────────────────────────────────────────
s = add_slide()
kicker(s, "4 · Vibecoding")
title(s, "18 pages construites, chacune vise une requête perdue")
rows = [
    ["Page", "Requête cible", "Armes GEO embarquées"],
    ["Le facialisme, c'est quoi ?", "définition de la catégorie", "DefinedTerm + FAQ : LA page de référence de l'entité"],
    ["Massage kobido à Paris", "kobido paris", "LocalBusiness, comparatif honnête avec prix concurrents"],
    ["Gua sha : guide complet", "gua sha utilisation", "HowTo 4 étapes, tableau d'erreurs, chiffres (15°, 5 min)"],
    ["Anti-âge sans injection", "botox naturel", "Honnêteté E-E-A-T (« rien ne remplace le botox, mais… »)"],
    ["Routine visage par saison", "routine skincare", "ADN de marque + mise à jour trimestrielle = fraîcheur"],
    ["Meilleur soin visage Paris", "meilleur soin paris", "ItemList, comparatif 5 formats, prix 30-180 €"],
    ["+ 12 articles magazine", "backlog Make publié", "Même gabarit GEO : answer box, stats, FAQ, JSON-LD"],
]
table(s, rows, 0.9, 2.3, 11.5, col_widths=[2.8, 2.3, 5.4], font_size=12, row_h=0.5)
tb = box(s, 0.9, 6.45, 11.5, 0.5)
para(tb, "+ hub index.html + llms.txt — réponse directe en 1er paragraphe, stats, FAQ, JSON-LD et date sur chaque page.",
     size=13, color=INK, first=True)
footer(s, 6)

# ── Slide 7 : Make ─────────────────────────────────────────────────────
s = add_slide()
kicker(s, "5 · Automatisation Make")
title(s, "Le robot éditorial : 1 article toutes les 30 minutes")
tb = box(s, 0.9, 2.15, 11.5, 0.6)
para(tb, "Schedule 30 min → Sheets (sujet « À générer ») → Claude (article JSON) → Google Doc → "
         "Sheets (maj ligne) → Gmail + Telegram", size=15, color=DEEP, bold=True, first=True)
rows = [
    ["Titre", "Résumé", "Mots clés", "Objectif", "Date de création", "Statut"],
    ["Bakuchiol vs rétinol : le match", "Généré par l'IA (3 phrases)", "bakuchiol, rétinol…", "GEO citation", "06/07/2026 14:30", "Généré"],
    ["Kobido à Bordeaux : le guide", "—", "—", "SEO local", "—", "À générer"],
]
table(s, rows, 0.9, 2.95, 11.5, col_widths=[2.6, 2.4, 1.8, 1.4, 1.7, 1.3], font_size=11, row_h=0.5)
stats = [
    ("48", "articles/jour possibles en démo — 1/jour recommandé en prod"),
    ("100 %", "human-in-the-loop : rien ne se publie sans relecture"),
    ("7", "patterns GEO encodés dans le system prompt"),
    ("~0,05 €", "coût IA par article (claude-sonnet-5)"),
]
stat_row(s, stats, y=4.85)
footer(s, 7)

# ── Slide 8 : KPIs ─────────────────────────────────────────────────────
s = add_slide()
kicker(s, "6 · Résultats attendus")
title(s, "KPIs : comment on saura que ça marche")
rows = [
    ["KPI", "Baseline (07/2026)", "Objectif 6 mois", "Mesure"],
    ["Part de voix IA (5 requêtes × 3 IA)", "~10 %", "60 %", "Grille de test mensuelle"],
    ["Citations tierces (listicles, presse)", "2", "10", "Veille mentions"],
    ["Trafic référent IA", "≈ 0", "500 sessions/mois", "GA4 (chatgpt.com, perplexity.ai)"],
    ["Positions Google R1-R5", "Non classé", "Top 10 sur 3/5", "Search Console"],
    ["Réservations studio via organique", "à poser", "+20 %", "UTM + module résa"],
]
table(s, rows, 0.9, 2.4, 11.5, col_widths=[3.6, 2.2, 2.2, 3.0], font_size=13, row_h=0.6)
footer(s, 8)

# ── Slide 9 : démo + conclusion ───────────────────────────────────────
s = add_slide()
kicker(s, "7 · Démo live")
title(s, "Ce qu'on vous montre maintenant")
demos = [
    ("1 · Le test IA en direct", "« Où faire un massage kobido à Paris ? » dans ChatGPT / Perplexity / Claude : qui est cité, quelles sources."),
    ("2 · Les 18 pages", "pages/index.html : structure, FAQ, JSON-LD visibles dans le code source."),
    ("3 · Le scénario Make", "Run once en direct : la ligne du Sheet passe « À générer » → « Généré », l'e-mail et le Telegram arrivent."),
]
for i, (h, txt) in enumerate(demos):
    c = card(s, 0.9, 2.25 + i * 1.0, 11.5, 0.9)
    tf = c.text_frame
    para(tf, h, size=14, color=DEEP, bold=True, font=SERIF, first=True, space_after=2)
    para(tf, txt, size=12, color=INK)
c = card(s, 0.9, 5.5, 11.5, 1.25, fill=DEEP)
tf = c.text_frame
para(tf, "Conclusion", size=13, color=PEACH, bold=True, first=True, space_after=3)
para(tf, "Le GEO n'est pas une discipline magique : c'est du SEO qui assume que le lecteur est parfois une "
         "machine. Structurer, chiffrer, dater, ouvrir ses crawlers — Seasonly a le produit et la légitimité ; "
         "il ne lui manquait que la couche citable. Elle existe désormais. Merci 🌿", size=13, color=CREAM)
footer(s, 9)

prs.save("pitch-deck-seasonly.pptx")
print("✓ pitch-deck-seasonly.pptx généré —", len(prs.slides.__iter__.__self__._sldIdLst), "slides")
